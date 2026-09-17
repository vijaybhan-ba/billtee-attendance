#!/usr/bin/env python3
"""Attendance 2026-27
====================
Python attendance system (Flask + OpenCV) with:
  * 8 employees seeded (Ajit ... Vijay)
  * face-capture check-in (webcam + OpenCV face detection)
  * fingerprint-code check-in (hook-up point for a real scanner)
  * 3-days-in-office / 3-days-outside rotation (two teams of 4, editable)
  * monthly reports with Excel (.xlsx) export
  * your GitHub link shown in the footer / reports (edit in Settings)
"""
import base64
import calendar
import json
import os
import sqlite3
from datetime import date, datetime
from io import BytesIO

import cv2
import numpy as np
from flask import (Flask, jsonify, render_template, request, send_file,
                   send_from_directory)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE, "attendance.db")
CONFIG_PATH = os.path.join(BASE, "config.json")
FACES_DIR = os.path.join(BASE, "faces")
CAPTURES_DIR = os.path.join(BASE, "captures")
BACKUPS_DIR = os.path.join(BASE, "backups")

CASCADE = cv2.CascadeClassifier(
    os.path.join(cv2.data.haarcascades, "haarcascade_frontalface_default.xml"))

DEFAULT_EMPLOYEES = ["Ajit", "Shyam", "Nikhat", "Nitish",
                     "Manish", "Anshuman", "Sarad", "Vijay"]

DEFAULT_CONFIG = {
    "github_link": "https://github.com/YOUR-USERNAME/attendance-2026-27",
    "cycle_start": "2026-04-01",   # day 1 of the 6-day rotation (3 in / 3 out)
    "year_label": "2026-27",
}

app = Flask(__name__)


# -------------------------------------------------------------------------- db
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    os.makedirs(FACES_DIR, exist_ok=True)
    os.makedirs(CAPTURES_DIR, exist_ok=True)
    os.makedirs(BACKUPS_DIR, exist_ok=True)
    conn = get_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS employees(
        id INTEGER PRIMARY KEY,
        name TEXT NOT NULL,
        grp TEXT NOT NULL DEFAULT 'A',
        fp_code TEXT,
        face_enrolled INTEGER NOT NULL DEFAULT 0)""")
    conn.execute("""CREATE TABLE IF NOT EXISTS attendance(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        employee_id INTEGER NOT NULL,
        day TEXT NOT NULL,
        checkin_time TEXT NOT NULL,
        method TEXT NOT NULL,
        location TEXT NOT NULL,
        capture TEXT,
        UNIQUE(employee_id, day))""")
    for i, name in enumerate(DEFAULT_EMPLOYEES, start=1):
        conn.execute(
            "INSERT OR IGNORE INTO employees(id, name, grp, fp_code) VALUES(?,?,?,?)",
            (i, name, "A" if i <= 4 else "B", str(1000 + i)))
    cols = [r["name"] for r in conn.execute("PRAGMA table_info(employees)")]
    if "office_loc" not in cols:
        conn.execute("ALTER TABLE employees ADD COLUMN office_loc TEXT")
        conn.execute("UPDATE employees SET office_loc = "
                     "CASE grp WHEN 'A' THEN 'Andheri' ELSE 'Ambadi' END")
    conn.commit()
    conn.close()


# ---------------------------------------------------------------------- config
def load_config():
    cfg = dict(DEFAULT_CONFIG)
    if os.path.exists(CONFIG_PATH):
        with open(CONFIG_PATH) as fh:
            cfg.update(json.load(fh))
    return cfg


def save_config(cfg):
    with open(CONFIG_PATH, "w") as fh:
        json.dump(cfg, fh, indent=2)


# ------------------------------------------------------------ rotation logic
def scheduled_location(grp, day):
    """3 days in office / 3 days outside. Team A starts in office on
    cycle_start, team B starts outside; they swap every 3 days."""
    cfg = load_config()
    start = date.fromisoformat(cfg["cycle_start"])
    pos = (day - start).days % 6
    a_in_office = pos < 3
    if grp == "A":
        return "Office" if a_in_office else "Outside"
    return "Outside" if a_in_office else "Office"


def duty_label(emp, day):
    loc = scheduled_location(emp["grp"], day)
    if loc == "Office":
        return f"Office – {emp['office_loc'] or 'Andheri'}"
    return "Outside"


# ---------------------------------------------------------------- face helpers
def decode_image(data_url):
    _, b64 = data_url.split(",", 1)
    arr = np.frombuffer(base64.b64decode(b64), dtype=np.uint8)
    return cv2.imdecode(arr, cv2.IMREAD_COLOR)


def count_faces(img):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    faces = CASCADE.detectMultiScale(gray, scaleFactor=1.1, minNeighbors=4,
                                     minSize=(50, 50))
    return len(faces)


# ----------------------------------------------------------------------- pages
@app.route("/")
def dashboard():
    cfg = load_config()
    today = date.today()
    conn = get_db()
    emps = conn.execute("SELECT * FROM employees ORDER BY id").fetchall()
    marks = {r["employee_id"]: r for r in conn.execute(
        "SELECT * FROM attendance WHERE day = ?", (today.isoformat(),))}
    conn.close()
    roster = []
    for e in emps:
        roster.append({
            "emp": e,
            "location": scheduled_location(e["grp"], today),
            "label": duty_label(e, today),
            "mark": marks.get(e["id"]),
        })
    pos = (today - date.fromisoformat(cfg["cycle_start"])).days % 6
    return render_template("dashboard.html", cfg=cfg, today=today,
                           roster=roster, cycle_day=pos + 1)


@app.route("/checkin")
def checkin_page():
    cfg = load_config()
    conn = get_db()
    emps = conn.execute("SELECT * FROM employees ORDER BY id").fetchall()
    done = conn.execute(
        """SELECT a.*, e.name FROM attendance a JOIN employees e ON e.id=a.employee_id
           WHERE a.day = ? ORDER BY a.checkin_time""",
        (date.today().isoformat(),)).fetchall()
    conn.close()
    return render_template("checkin.html", cfg=cfg, emps=emps, done=done,
                           today=date.today())


@app.route("/employees")
def employees_page():
    cfg = load_config()
    conn = get_db()
    emps = conn.execute("SELECT * FROM employees ORDER BY id").fetchall()
    conn.close()
    return render_template("employees.html", cfg=cfg, emps=emps)


@app.route("/reports")
def reports_page():
    cfg = load_config()
    month = request.args.get("month", date.today().strftime("%Y-%m"))
    grid, emps = build_grid(month)
    backups = sorted((f for f in os.listdir(BACKUPS_DIR) if f.endswith(".json")),
                     reverse=True) if os.path.isdir(BACKUPS_DIR) else []
    return render_template("reports.html", cfg=cfg, month=month, grid=grid,
                           emps=emps, month_name=month_title(month),
                           prev_month=shift_month(month, -1), backups=backups)


@app.route("/settings", methods=["GET", "POST"])
def settings_page():
    cfg = load_config()
    msg = None
    if request.method == "POST":
        cfg["github_link"] = request.form.get("github_link", "").strip() or cfg["github_link"]
        cfg["cycle_start"] = request.form.get("cycle_start", "").strip() or cfg["cycle_start"]
        save_config(cfg)
        msg = "Settings saved."
    return render_template("settings.html", cfg=cfg, msg=msg)


# ------------------------------------------------------------------------ api
@app.route("/api/duty")
def api_duty():
    conn = get_db()
    emp = conn.execute("SELECT * FROM employees WHERE id = ?",
                       (request.args.get("emp", 0),)).fetchone()
    conn.close()
    if not emp:
        return jsonify(error="unknown employee"), 404
    return jsonify(location=scheduled_location(emp["grp"], date.today()),
                   label=duty_label(emp, date.today()))


@app.route("/api/checkin", methods=["POST"])
def api_checkin():
    data = request.get_json(force=True)
    emp_id = int(data.get("employee_id", 0))
    method = data.get("method")
    conn = get_db()
    emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
    if not emp:
        conn.close()
        return jsonify(error="Unknown employee"), 404

    location = scheduled_location(emp["grp"], date.today())
    label = (f"Office – {emp['office_loc'] or 'Andheri'}"
             if location == "Office" else "Outside")
    capture = None

    if method == "face":
        if not data.get("image"):
            conn.close()
            return jsonify(error="No camera image received. Start the camera first."), 400
        img = decode_image(data["image"])
        if img is None:
            conn.close()
            return jsonify(error="Could not read the camera image."), 400
        if count_faces(img) == 0:
            conn.close()
            return jsonify(error="No face detected in the capture. Face the camera and try again."), 400
        fname = f"{emp_id}_{date.today().isoformat()}_{datetime.now().strftime('%H%M%S')}.jpg"
        cv2.imwrite(os.path.join(CAPTURES_DIR, fname), img)
        capture = fname
    elif method == "fingerprint":
        if not emp["fp_code"] or data.get("fp_code", "").strip() != emp["fp_code"]:
            conn.close()
            return jsonify(error="Fingerprint code does not match."), 403
    elif method == "manual":
        pass
    else:
        conn.close()
        return jsonify(error="Unknown method. Use face, fingerprint or manual."), 400

    now = datetime.now()
    conn.execute(
        """INSERT INTO attendance(employee_id, day, checkin_time, method, location, capture)
           VALUES(?,?,?,?,?,?)
           ON CONFLICT(employee_id, day) DO UPDATE SET
             checkin_time=excluded.checkin_time, method=excluded.method,
             location=excluded.location, capture=excluded.capture""",
        (emp_id, date.today().isoformat(), now.strftime("%H:%M:%S"),
         method, location, capture))
    conn.commit()
    conn.close()
    return jsonify(ok=True, name=emp["name"], location=location, label=label,
                   time=now.strftime("%H:%M:%S"), method=method)


@app.route("/api/enroll", methods=["POST"])
def api_enroll():
    data = request.get_json(force=True)
    emp_id = int(data.get("employee_id", 0))
    conn = get_db()
    emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
    if not emp:
        conn.close()
        return jsonify(error="Unknown employee"), 404
    if not data.get("image"):
        conn.close()
        return jsonify(error="No camera image received."), 400
    img = decode_image(data["image"])
    if img is None or count_faces(img) == 0:
        conn.close()
        return jsonify(error="No face detected. Face the camera and try again."), 400
    cv2.imwrite(os.path.join(FACES_DIR, f"{emp_id}.jpg"), img)
    conn.execute("UPDATE employees SET face_enrolled = 1 WHERE id = ?", (emp_id,))
    conn.commit()
    conn.close()
    return jsonify(ok=True, name=emp["name"])


@app.route("/api/employee", methods=["POST"])
def api_employee():
    data = request.get_json(force=True)
    emp_id = int(data.get("employee_id", 0))
    grp = data.get("grp")
    fp_code = data.get("fp_code")
    office_loc = data.get("office_loc")
    if grp not in ("A", "B"):
        return jsonify(error="Group must be A or B"), 400
    if office_loc is not None and office_loc not in ("Andheri", "Ambadi"):
        return jsonify(error="Office location must be Andheri or Ambadi"), 400
    conn = get_db()
    emp = conn.execute("SELECT * FROM employees WHERE id = ?", (emp_id,)).fetchone()
    if not emp:
        conn.close()
        return jsonify(error="Unknown employee"), 404
    conn.execute(
        "UPDATE employees SET grp = ?, fp_code = ?, office_loc = ? WHERE id = ?",
        (grp,
         str(fp_code).strip() if fp_code is not None else emp["fp_code"],
         office_loc if office_loc is not None else emp["office_loc"],
         emp_id))
    conn.commit()
    conn.close()
    return jsonify(ok=True)


# --------------------------------------------------------------------- report
def month_title(month):
    y, m = (int(x) for x in month.split("-"))
    return f"{calendar.month_name[m]} {y}"


def shift_month(month, delta):
    y, m = (int(x) for x in month.split("-"))
    m += delta
    y += (m - 1) // 12
    m = (m - 1) % 12 + 1
    return f"{y:04d}-{m:02d}"


def build_grid(month):
    y, m = (int(x) for x in month.split("-"))
    ndays = calendar.monthrange(y, m)[1]
    conn = get_db()
    emps = conn.execute("SELECT * FROM employees ORDER BY id").fetchall()
    rows = conn.execute(
        "SELECT employee_id, day, method, location, checkin_time FROM attendance")
    mark = {(r["employee_id"], r["day"]): r for r in rows}
    conn.close()
    grid = []
    for e in emps:
        cells = []
        counts = {"Present": 0, "Office": 0, "Outside": 0, "Absent": 0}
        for d in range(1, ndays + 1):
            key = (e["id"], f"{month}-{d:02d}")
            r = mark.get(key)
            if r:
                counts["Present"] += 1
                counts[r["location"]] += 1
                if r["location"] == "Office":
                    v = "AND" if (e["office_loc"] or "Andheri") == "Andheri" else "AMB"
                else:
                    v = "OUT"
                cells.append({"v": v,
                              "cls": "in" if r["location"] == "Office" else "out",
                              "title": f"{r['method']} @ {r['checkin_time']}"})
            else:
                counts["Absent"] += 1
                cells.append({"v": "—", "cls": "abs", "title": ""})
        grid.append({"emp": e, "cells": cells, "counts": counts})
    return grid, emps


def add_month_sheet(wb, month, cfg):
    """One styled worksheet per month."""
    y, m = (int(x) for x in month.split("-"))
    ndays = calendar.monthrange(y, m)[1]
    grid, emps = build_grid(month)

    ws = wb.create_sheet(title=f"{calendar.month_abbr[m]}-{y}")
    ws["A1"] = f"Attendance {cfg['year_label']} — {calendar.month_name[m]} {y}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = cfg["github_link"]
    ws["A2"].font = Font(italic=True, color="666666")

    header = ["No", "Employee", "Team", "Office Location"] + \
             [str(d) for d in range(1, ndays + 1)] + \
             ["Present", "Office", "Outside", "Absent"]
    ws.append([])
    ws.append(header)
    hdr_fill = PatternFill("solid", fgColor="2F5496")
    for col in range(1, len(header) + 1):
        c = ws.cell(row=4, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill

    fills = {"in": PatternFill("solid", fgColor="C6EFCE"),
             "out": PatternFill("solid", fgColor="BDD7EE"),
             "abs": PatternFill("solid", fgColor="F2DCDB")}
    for g in grid:
        row = [g["emp"]["id"], g["emp"]["name"], g["emp"]["grp"],
               g["emp"]["office_loc"] or "Andheri"]
        row += [c["v"] for c in g["cells"]]
        row += [g["counts"]["Present"], g["counts"]["Office"],
                g["counts"]["Outside"], g["counts"]["Absent"]]
        ws.append(row)
        r = ws.max_row
        for i, c in enumerate(g["cells"], start=5):
            ws.cell(row=r, column=i).fill = fills[c["cls"]]
    ws.column_dimensions["B"].width = 14
    return ws


@app.route("/api/report/xlsx")
def report_xlsx():
    cfg = load_config()
    month = request.args.get("month", date.today().strftime("%Y-%m"))
    two = request.args.get("months") == "2"
    months = [shift_month(month, -1), month] if two else [month]

    wb = Workbook()
    wb.remove(wb.active)
    for mo in months:
        add_month_sheet(wb, mo, cfg)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = (f"attendance_{months[0]}_to_{months[-1]}.xlsx" if two
            else f"attendance_{month}.xlsx")
    return send_file(
        buf, as_attachment=True, download_name=name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/api/backup")
def backup():
    """2-month backup (selected month + previous): JSON with employees and
    attendance; a copy is kept on disk in backups/."""
    cfg = load_config()
    month = request.args.get("month", date.today().strftime("%Y-%m"))
    months = [shift_month(month, -1), month]
    conn = get_db()
    employees = [dict(r) for r in conn.execute(
        "SELECT * FROM employees ORDER BY id")]
    att = [dict(r) for r in conn.execute(
        "SELECT * FROM attendance WHERE substr(day,1,7) IN (?,?) "
        "ORDER BY day, employee_id", months)]
    conn.close()
    payload = {
        "type": "attendance-backup",
        "year_label": cfg["year_label"],
        "github_link": cfg["github_link"],
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "months": months,
        "employees": employees,
        "attendance": att,
    }
    fname = f"backup_{date.today().isoformat()}_{month}.json"
    with open(os.path.join(BACKUPS_DIR, fname), "w") as fh:
        json.dump(payload, fh, indent=2)
    return send_file(os.path.join(BACKUPS_DIR, fname), as_attachment=True,
                     download_name=fname, mimetype="application/json")


@app.route("/api/backup/download")
def backup_download():
    name = os.path.basename(request.args.get("name", ""))
    if not name.startswith("backup_") or not name.endswith(".json"):
        return jsonify(error="Not a backup file"), 400
    return send_from_directory(BACKUPS_DIR, name, as_attachment=True)


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=False)
